import base64
import io
import logging
from pathlib import Path
from typing import Optional

import PyPDF2
import docx
import pandas as pd

try:
    from pdf2image import convert_from_path
    PDF_SUPPORT = True
except Exception:
    PDF_SUPPORT = False

try:
    import pytesseract
    OCR_SUPPORT = True
    from PIL import Image
except Exception:
    OCR_SUPPORT = False
    Image = None

try:
    import openpyxl
    EXCEL_SUPPORT = True
except Exception:
    EXCEL_SUPPORT = False
    openpyxl = None

logger = logging.getLogger(__name__)


class FileProcessor:
    """Utility class to load various document formats."""

    @staticmethod
    def _read_pdf(path: Path) -> str:
        text = ""
        try:
            with path.open("rb") as f:
                reader = PyPDF2.PdfReader(f)
                for idx, page in enumerate(reader.pages):
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
                    elif PDF_SUPPORT and OCR_SUPPORT and Image:
                        images = convert_from_path(str(path), first_page=idx + 1, last_page=idx + 1)
                        if images:
                            ocr = pytesseract.image_to_string(images[0], lang="jpn+eng")
                            if ocr.strip():
                                text += ocr + "\n"
        except Exception as e:  # pragma: no cover - parsing errors
            logger.error("PDF read error: %s", e)
        return text

    @staticmethod
    def _read_docx(path: Path) -> str:
        text = ""
        try:
            doc = docx.Document(str(path))
            for para in doc.paragraphs:
                text += para.text + "\n"
            if OCR_SUPPORT and Image:
                for rel in doc.part.related_parts.values():
                    if "image" in rel.content_type:
                        img = Image.open(io.BytesIO(rel.blob))
                        ocr = pytesseract.image_to_string(img, lang="jpn+eng")
                        if ocr.strip():
                            text += ocr + "\n"
        except Exception as e:  # pragma: no cover - parsing errors
            logger.error("DOCX read error: %s", e)
        return text

    @staticmethod
    def _read_excel(path: Path) -> str:
        text = ""
        try:
            if EXCEL_SUPPORT:
                wb = openpyxl.load_workbook(path, data_only=True)
                for sheet in wb.worksheets:
                    text += f"# {sheet.title}\n"
                    for row in sheet.iter_rows(values_only=True):
                        cells = ["" if c is None else str(c) for c in row]
                        text += "\t".join(cells) + "\n"
                    if OCR_SUPPORT and Image and getattr(sheet, "_images", []):
                        for img in sheet._images:
                            try:
                                img_bytes = img._data()
                                img_obj = Image.open(io.BytesIO(img_bytes))
                                ocr = pytesseract.image_to_string(img_obj, lang="jpn+eng")
                                if ocr.strip():
                                    text += ocr + "\n"
                            except Exception:
                                pass
            else:
                df = pd.read_excel(path, sheet_name=None)
                for name, sheet in df.items():
                    text += f"# {name}\n"
                    text += sheet.to_string() + "\n"
        except Exception as e:  # pragma: no cover - parsing errors
            logger.error("Excel read error: %s", e)
        return text

    @staticmethod
    def _read_text(path: Path) -> str:
        try:
            return path.read_text(encoding="utf-8", errors="replace")
        except Exception as e:  # pragma: no cover
            logger.error("Text read error: %s", e)
            return ""

    @classmethod
    def process_file(cls, file_path: Path) -> str:
        ext = file_path.suffix.lower()
        if ext == ".pdf":
            return cls._read_pdf(file_path)
        if ext == ".docx":
            return cls._read_docx(file_path)
        if ext in {".xls", ".xlsx"}:
            return cls._read_excel(file_path)
        return cls._read_text(file_path)

